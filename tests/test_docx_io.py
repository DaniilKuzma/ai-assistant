from pathlib import Path

from docx import Document

from src.docx.docx_corrector import count_docx_text_blocks, correct_docx, correct_docx_incremental
from src.runtime.edit_realizer import apply_runtime_edits
from src.schema.edits import CorrectionResult, RuntimeEdit


FAKE_EDITS = [
    RuntimeEdit(2, 8, "незнаю", "не знаю", "split_join", "fake_split", 0.99, "Раздельное написание с не."),
    RuntimeEdit(15, 15, "", ",", "punctuation", "fake_comma", 0.98, "Запятая перед придаточной частью."),
]
SAFE_FAKE_EDITS = [FAKE_EDITS[0]]


class FakeCorrector:
    def __init__(self, corrections: dict[str, str], *, falsey: bool = False) -> None:
        self.corrections = corrections
        self.falsey = falsey
        self.calls: list[str] = []

    def __bool__(self) -> bool:
        return not self.falsey

    def correct(self, text: str) -> CorrectionResult:
        self.calls.append(text)
        corrected_text = self.corrections.get(text, text)
        edits = _fake_edits_for_text(text) if corrected_text != text else []
        return CorrectionResult(text, corrected_text, edits)


class ExplicitEditCorrector:
    def __init__(self, edits_by_text: dict[str, list[RuntimeEdit]]) -> None:
        self.edits_by_text = edits_by_text
        self.calls: list[str] = []

    def correct(self, text: str) -> CorrectionResult:
        self.calls.append(text)
        edits = list(self.edits_by_text.get(text, []))
        return CorrectionResult(text, apply_runtime_edits(text, edits), edits)


def _fake_edits_for_text(text: str) -> list[RuntimeEdit]:
    source = "\u043d\u0435\u0437\u043d\u0430\u044e"
    replacement = "\u043d\u0435 \u0437\u043d\u0430\u044e"
    if source not in text:
        return []
    start = text.index(source)
    return [
        RuntimeEdit(
            start,
            start + len(source),
            source,
            replacement,
            "split_join",
            "fake_split",
            0.99,
            "Раздельное написание с не.",
        )
    ]


def test_docx_correction_preserves_paragraph_count(tmp_path: Path):
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "output.docx"

    document = Document()
    document.add_paragraph("Я незнаю что делать")
    document.add_paragraph("Чистый текст.")
    document.save(input_path)

    fake_corrector = FakeCorrector({"Я незнаю что делать": "Я не знаю, что делать."})
    edits = correct_docx(input_path, output_path, corrector=fake_corrector)

    corrected = Document(output_path)
    paragraphs = [paragraph.text for paragraph in corrected.paragraphs]

    assert paragraphs == ["Я не знаю что делать", "Чистый текст."]
    assert fake_corrector.calls == ["Я незнаю что делать", "Чистый текст."]
    assert edits == SAFE_FAKE_EDITS


def test_docx_correction_writes_runtime_corrected_paragraph_for_tiny_docx(tmp_path: Path):
    input_path = tmp_path / "tiny.docx"
    output_path = tmp_path / "tiny_output.docx"

    document = Document()
    document.add_paragraph("Он незнал что делать")
    document.save(input_path)

    edit = RuntimeEdit(
        start=3,
        end=9,
        source="незнал",
        replacement="не знал",
        edit_type="split_join",
        rule_id="ne_verb",
        confidence=1.0,
        explanation="Частица не с глаголом пишется раздельно.",
    )
    fake_corrector = FakeCorrector({"Он незнал что делать": "Он не знал что делать"})
    fake_corrector.correct = lambda text: CorrectionResult(text, "Он не знал что делать", [edit])

    edits = correct_docx(input_path, output_path, corrector=fake_corrector)

    corrected = Document(output_path)
    assert output_path.exists()
    assert corrected.paragraphs[0].text == "Он не знал что делать"
    assert edits == [edit]


def test_docx_correction_preserves_first_run_bold_formatting(tmp_path: Path):
    input_path = tmp_path / "styled.docx"
    output_path = tmp_path / "styled_output.docx"

    document = Document()
    run = document.add_paragraph().add_run("Я незнаю что делать")
    run.bold = True
    document.save(input_path)

    fake_corrector = FakeCorrector({"Я незнаю что делать": "Я не знаю, что делать."})
    edits = correct_docx(input_path, output_path, corrector=fake_corrector)

    corrected = Document(output_path)
    assert corrected.paragraphs[0].runs[0].text == "Я не знаю что делать"
    assert corrected.paragraphs[0].runs[0].bold is True
    assert fake_corrector.calls == ["Я незнаю что делать"]
    assert edits == SAFE_FAKE_EDITS


def test_docx_correction_preserves_inline_run_formatting_boundaries(tmp_path: Path):
    input_path = tmp_path / "inline_styled.docx"
    output_path = tmp_path / "inline_styled_output.docx"

    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Я ")
    bold_run = paragraph.add_run("незнаю")
    bold_run.bold = True
    italic_run = paragraph.add_run(" что делать")
    italic_run.italic = True
    document.save(input_path)

    fake_corrector = FakeCorrector({"Я незнаю что делать": "Я не знаю, что делать."})
    edits = correct_docx(input_path, output_path, corrector=fake_corrector)

    corrected = Document(output_path)
    paragraph = corrected.paragraphs[0]
    assert paragraph.text == "Я не знаю что делать"
    assert paragraph.runs[0].text == "Я "
    assert paragraph.runs[1].text == "не знаю"
    assert paragraph.runs[1].bold is True
    assert paragraph.runs[2].text == " что делать"
    assert paragraph.runs[2].italic is True
    assert fake_corrector.calls == ["Я незнаю что делать"]
    assert edits == SAFE_FAKE_EDITS


def test_docx_correction_uses_injected_corrector_without_constructing_default(tmp_path: Path, monkeypatch):
    input_path = tmp_path / "injected.docx"
    output_path = tmp_path / "injected_output.docx"

    document = Document()
    document.add_paragraph("Я незнаю что делать")
    document.save(input_path)

    def fail_if_default_corrector_is_constructed():
        raise AssertionError("DOCX correction must use the injected corrector")

    monkeypatch.setattr("src.docx.docx_corrector.Corrector", fail_if_default_corrector_is_constructed)
    fake_corrector = FakeCorrector({"Я незнаю что делать": "Я не знаю, что делать."}, falsey=True)

    edits = correct_docx(input_path, output_path, corrector=fake_corrector)

    corrected = Document(output_path)
    assert corrected.paragraphs[0].text == "Я не знаю что делать"
    assert fake_corrector.calls == ["Я незнаю что делать"]
    assert edits == SAFE_FAKE_EDITS


def test_incremental_docx_first_run_checks_non_empty_paragraphs_and_writes_output(tmp_path: Path):
    input_path = tmp_path / "incremental_first.docx"
    output_path = tmp_path / "incremental_first_output.docx"

    document = Document()
    document.add_paragraph("Я незнаю что делать")
    document.add_paragraph("")
    document.add_paragraph("   ")
    document.add_paragraph("Чистый текст.")
    document.save(input_path)

    fake_corrector = FakeCorrector({"Я незнаю что делать": "Я не знаю, что делать."})
    result = correct_docx_incremental(input_path, output_path, corrector=fake_corrector)

    corrected = Document(output_path)
    paragraphs = [paragraph.text for paragraph in corrected.paragraphs]

    assert output_path.exists()
    assert paragraphs == ["Я не знаю что делать", "", "   ", "Чистый текст."]
    assert fake_corrector.calls == ["Я незнаю что делать", "Чистый текст."]
    assert result.checked_paragraphs == 2
    assert result.reused_paragraphs == 0
    assert len(result.paragraph_cache) == 2
    assert result.edits == SAFE_FAKE_EDITS


def test_incremental_docx_second_run_reuses_same_document_paragraphs(tmp_path: Path):
    input_path = tmp_path / "incremental_reuse.docx"
    first_output_path = tmp_path / "incremental_reuse_first_output.docx"
    second_output_path = tmp_path / "incremental_reuse_second_output.docx"

    document = Document()
    document.add_paragraph("Я незнаю что делать")
    document.add_paragraph("")
    document.add_paragraph("Чистый текст.")
    document.save(input_path)

    first_corrector = FakeCorrector({"Я незнаю что делать": "Я не знаю, что делать."})
    first = correct_docx_incremental(input_path, first_output_path, corrector=first_corrector)

    second_corrector = FakeCorrector({})
    second = correct_docx_incremental(
        input_path,
        second_output_path,
        corrector=second_corrector,
        previous_cache=first.paragraph_cache,
    )

    corrected = Document(second_output_path)
    paragraphs = [paragraph.text for paragraph in corrected.paragraphs]

    assert second_output_path.exists()
    assert paragraphs == ["Я не знаю что делать", "", "Чистый текст."]
    assert second_corrector.calls == []
    assert second.checked_paragraphs == 0
    assert second.reused_paragraphs == 2
    assert len(second.paragraph_cache) == 2
    assert second.edits == SAFE_FAKE_EDITS


def test_docx_correction_filters_final_punctuation_but_keeps_controlled_casing(tmp_path: Path):
    input_path = tmp_path / "structural.docx"
    output_path = tmp_path / "structural_output.docx"
    source = "\u0420\u043e\u0441\u0441\u0438\u0439\u0441\u043a\u043e\u0439 \u0444\u0435\u0434\u0435\u0440\u0430\u0446\u0438\u0438"
    expected = "\u0420\u043e\u0441\u0441\u0438\u0439\u0441\u043a\u043e\u0439 \u0424\u0435\u0434\u0435\u0440\u0430\u0446\u0438\u0438"

    document = Document()
    document.add_paragraph(source)
    document.save(input_path)

    final_dot = RuntimeEdit(len(source), len(source), "", ".", "punctuation", "final_punctuation", 1.0)
    edits = correct_docx(input_path, output_path, corrector=ExplicitEditCorrector({source: [final_dot]}))

    corrected = Document(output_path)
    assert corrected.paragraphs[0].text == expected
    assert [(edit.source, edit.replacement, edit.rule_id) for edit in edits] == [
        ("\u0444\u0435\u0434\u0435\u0440\u0430\u0446\u0438\u0438", "\u0424\u0435\u0434\u0435\u0440\u0430\u0446\u0438\u0438", "casing_geo_names")
    ]


def test_docx_correction_skips_edits_inside_quoted_examples(tmp_path: Path):
    input_path = tmp_path / "quoted_example.docx"
    output_path = tmp_path / "quoted_example_output.docx"
    source = (
        "\u041d\u0430\u043f\u0440\u0438\u043c\u0435\u0440, \u00ab\u0421\u0442\u0443\u0434\u0435\u043d\u0442 "
        "\u043d\u0435\u0437\u043d\u0430\u043b \u043e\u0442\u0432\u0435\u0442\u0430\u00bb."
    )
    wrong = "\u043d\u0435\u0437\u043d\u0430\u043b"
    replacement = "\u043d\u0435 \u0437\u043d\u0430\u043b"

    document = Document()
    document.add_paragraph(source)
    document.save(input_path)

    edit_start = source.index(wrong)
    edit = RuntimeEdit(edit_start, edit_start + len(wrong), wrong, replacement, "split_join", "ne_verb", 1.0)
    edits = correct_docx(input_path, output_path, corrector=ExplicitEditCorrector({source: [edit]}))

    corrected = Document(output_path)
    assert corrected.paragraphs[0].text == source
    assert edits == []


def test_docx_correction_skips_context_dependent_service_word_edits(tmp_path: Path):
    input_path = tmp_path / "service_word.docx"
    output_path = tmp_path / "service_word_output.docx"
    source = (
        "\u042d\u0442\u043e \u043e\u0434\u043d\u043e \u0438 \u0442\u043e \u0436\u0435 "
        "\u0441\u043b\u043e\u0432\u043e."
    )
    wrong = "\u0442\u043e \u0436\u0435"

    document = Document()
    document.add_paragraph(source)
    document.save(input_path)

    edit_start = source.index(wrong)
    edit = RuntimeEdit(
        edit_start,
        edit_start + len(wrong),
        wrong,
        "\u0442\u043e\u0436\u0435",
        "split_join",
        "compound_service_words",
        1.0,
    )
    edits = correct_docx(input_path, output_path, corrector=ExplicitEditCorrector({source: [edit]}))

    corrected = Document(output_path)
    assert corrected.paragraphs[0].text == source
    assert edits == []


def test_docx_correction_checks_table_cells_without_adding_final_punctuation(tmp_path: Path):
    input_path = tmp_path / "table.docx"
    output_path = tmp_path / "table_output.docx"
    source = "\u041a\u0442\u043e \u0442\u043e"
    replacement = "\u041a\u0442\u043e-\u0442\u043e"

    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = source
    document.save(input_path)

    hyphen = RuntimeEdit(0, len(source), source, replacement, "hyphen", "hyphen_particles", 1.0)
    final_dot = RuntimeEdit(len(source), len(source), "", ".", "punctuation", "final_punctuation", 1.0)
    edits = correct_docx(input_path, output_path, corrector=ExplicitEditCorrector({source: [hyphen, final_dot]}))

    corrected = Document(output_path)
    assert corrected.tables[0].cell(0, 0).text == replacement
    assert [(edit.source, edit.replacement, edit.rule_id) for edit in edits] == [(source, replacement, "hyphen_particles")]


def test_count_docx_text_blocks_includes_table_cells(tmp_path: Path):
    input_path = tmp_path / "count_blocks.docx"

    document = Document()
    document.add_paragraph("\u041e\u0441\u043d\u043e\u0432\u043d\u043e\u0439 \u0430\u0431\u0437\u0430\u0446")
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "\u042f\u0447\u0435\u0439\u043a\u0430"
    table.cell(1, 0).text = "   "
    document.save(input_path)

    assert count_docx_text_blocks(input_path) == 2


def test_incremental_docx_changed_paragraph_checks_only_changed_paragraph(tmp_path: Path):
    first_input_path = tmp_path / "incremental_changed_first.docx"
    second_input_path = tmp_path / "incremental_changed_second.docx"
    first_output_path = tmp_path / "incremental_changed_first_output.docx"
    second_output_path = tmp_path / "incremental_changed_second_output.docx"

    first_document = Document()
    first_document.add_paragraph("Первый незнаю")
    first_document.add_paragraph("Второй текст.")
    first_document.save(first_input_path)

    second_document = Document()
    second_document.add_paragraph("Первый незнаю")
    second_document.add_paragraph("Второй незнаю")
    second_document.save(second_input_path)

    first_corrector = FakeCorrector(
        {
            "Первый незнаю": "Первый не знаю",
            "Второй текст.": "Второй текст.",
        }
    )
    first = correct_docx_incremental(first_input_path, first_output_path, corrector=first_corrector)

    second_corrector = FakeCorrector({"Второй незнаю": "Второй не знаю"})
    second = correct_docx_incremental(
        second_input_path,
        second_output_path,
        corrector=second_corrector,
        previous_cache=first.paragraph_cache,
    )

    corrected = Document(second_output_path)
    paragraphs = [paragraph.text for paragraph in corrected.paragraphs]

    assert second_output_path.exists()
    assert paragraphs == ["Первый не знаю", "Второй не знаю"]
    assert second_corrector.calls == ["Второй незнаю"]
    assert second.checked_paragraphs == 1
    assert second.reused_paragraphs == 1
    assert len(second.paragraph_cache) == 2
