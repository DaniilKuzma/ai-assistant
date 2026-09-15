from __future__ import annotations

from dataclasses import dataclass
import difflib
import html
from pathlib import Path
import re
import tempfile
from typing import Any

from docx import Document

from src.config.load_config import load_config
from src.docx.docx_corrector import correct_docx, correct_docx_incremental
from src.inference.incremental_corrector import IncrementalCorrector
from src.runtime.corrector import Corrector
from src.runtime.explanations import explanation_for
from src.schema.edits import RuntimeEdit


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CONFIG_PATH = PROJECT_ROOT / "configs" / "config.yaml"
st: Any | None = None


@dataclass(frozen=True)
class StreamlitCorrectorLoadResult:
    corrector: Any
    kind: str
    error: str | None = None


@dataclass(frozen=True)
class DiffTooltip:
    source: str
    replacement: str
    text: str


def build_streamlit_corrector(config_path: str | Path = DEFAULT_CONFIG_PATH) -> StreamlitCorrectorLoadResult:
    config = load_config(Path(config_path))
    corrector = Corrector.from_config(config)
    kind = "direct_neural" if getattr(corrector, "neural_backend", None) is not None else "deterministic_fallback"
    return StreamlitCorrectorLoadResult(corrector, kind)


def correct_text_for_gui(
    corrector: Any,
    source: str,
    *,
    use_incremental: bool,
    previous_cache: dict[str, Any] | None = None,
) -> Any:
    incremental_corrector = IncrementalCorrector(corrector)
    if use_incremental:
        return incremental_corrector.correct_incremental(source, previous_cache)
    return incremental_corrector.correct_incremental(source)


def _cached_streamlit_corrector(config_path: str = str(DEFAULT_CONFIG_PATH)) -> StreamlitCorrectorLoadResult:
    return build_streamlit_corrector(config_path)


def main() -> None:
    global st
    st = _import_streamlit()
    cached_streamlit_corrector = st.cache_resource(show_spinner="Загрузка модели корректора...")(
        _cached_streamlit_corrector
    )
    st.set_page_config(page_title="Система исправления ошибок русского языка", layout="wide")
    st.title("Система исправления ошибок русского языка")

    tab_text, tab_docx = st.tabs(["Текст", "DOCX-документ"])

    with tab_text:
        load_result = cached_streamlit_corrector(str(DEFAULT_CONFIG_PATH))
        corrector = load_result.corrector
        use_incremental = st.checkbox("Повторно проверять только изменённые фрагменты", value=False)
        source = st.text_area("Исходный текст", height=180)
        if st.button("Исправить текст", type="primary") and source.strip():
            cache_key = text_incremental_cache_key()
            previous_text_segment_cache = st.session_state.setdefault("previous_text_segment_cache", {})
            last_source_text = st.session_state.setdefault("last_source_text", {})
            last_corrected_text = st.session_state.setdefault("last_corrected_text", {})
            incremental_summary = None
            previous_cache = previous_text_segment_cache.get(cache_key) if use_incremental else None
            result = correct_text_for_gui(
                corrector,
                source,
                use_incremental=use_incremental,
                previous_cache=previous_cache,
            )
            previous_text_segment_cache[cache_key] = result.segment_cache
            if use_incremental:
                incremental_summary = build_incremental_summary(result)
            last_source_text[cache_key] = source
            last_corrected_text[cache_key] = result.corrected_text
            latest_text_correction = {
                "source": source,
                "result": result,
            }
            if incremental_summary is not None:
                latest_text_correction["incremental_summary"] = incremental_summary
            st.session_state["latest_text_correction"] = latest_text_correction
        _render_latest_text_correction()

    with tab_docx:
        uploaded = st.file_uploader("Word-документ", type=["docx"])
        use_docx_incremental = st.checkbox("Повторно проверять только изменённые абзацы", value=False)
        if uploaded is not None:
            if docx_correction_requested(uploaded.name):
                docx_load_result = cached_streamlit_corrector(str(DEFAULT_CONFIG_PATH))
                corrector = docx_load_result.corrector
                with tempfile.TemporaryDirectory() as tmp:
                    input_path = Path(tmp) / "input.docx"
                    output_path = Path(tmp) / "corrected.docx"
                    input_path.write_bytes(uploaded.getvalue())
                    docx_incremental_summary = None
                    if use_docx_incremental:
                        cache_key = docx_cache_key(uploaded.name)
                        paragraph_cache_by_upload = st.session_state.setdefault("docx_paragraph_cache_by_upload", {})
                        result = correct_docx_incremental(
                            input_path,
                            output_path,
                            corrector,
                            previous_cache=paragraph_cache_by_upload.get(cache_key),
                        )
                        paragraph_cache_by_upload[cache_key] = result.paragraph_cache
                        edits = result.edits
                        docx_incremental_summary = build_docx_incremental_summary(result)
                    else:
                        edits = correct_docx(input_path, output_path, corrector)
                    st.download_button(
                        "Скачать исправленный DOCX",
                        data=output_path.read_bytes(),
                        file_name="corrected.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                    render_docx_incremental_status(
                        docx_incremental_summary,
                        enabled=use_docx_incremental,
                    )
                    render_edit_panel(edits)


def build_edit_rows(edits: list[RuntimeEdit]) -> list[dict[str, object]]:
    return [
        {
            "Было": _user_facing_edit_value(edit.source, empty_text="Отсутствует"),
            "Стало": _user_facing_edit_value(edit.replacement, empty_text="Удалить"),
            "Правило": _user_facing_rule_text(edit),
        }
        for edit in edits
    ]


def build_edit_summary(edits: list[RuntimeEdit]) -> dict[str, int]:
    punctuation_count = sum(1 for edit in edits if str(getattr(edit, "edit_type", "") or "") == "punctuation")
    total = len(edits)
    return {
        "total": total,
        "orthography": total - punctuation_count,
        "punctuation": punctuation_count,
    }


def render_edit_panel(edits: list[RuntimeEdit]) -> None:
    rows = build_edit_rows(edits)

    st.markdown("### Правки")
    render_edit_summary(edits)

    if not rows:
        st.info("Правок нет")
        return

    for index, row in enumerate(rows):
        with st.expander(_edit_card_title(row), expanded=index == 0):
            st.markdown(f"**Было:** {row['Было']}")
            st.markdown(f"**Стало:** {row['Стало']}")
            st.markdown(f"**Правило:** {row['Правило']}")


def render_edit_summary(edits: list[RuntimeEdit]) -> None:
    summary = build_edit_summary(edits)
    st.markdown(_EDIT_PANEL_STYLE, unsafe_allow_html=True)
    st.markdown(
        (
            '<div class="edit-summary">'
            f"<span>Исправлений: {summary['total']}</span>"
            f"<span>Орфография: {summary['orthography']}</span>"
            f"<span>Пунктуация: {summary['punctuation']}</span>"
            "</div>"
        ),
        unsafe_allow_html=True,
    )


def _edit_card_title(row: dict[str, object]) -> str:
    return f"{row['Было']} → {row['Стало']}"


def _user_facing_edit_value(value: object, *, empty_text: str) -> str:
    text = str(value or "").strip()
    return text or empty_text


def _user_facing_rule_text(edit: RuntimeEdit) -> str:
    return (
        str(getattr(edit, "explanation", "") or "").strip()
        or explanation_for(str(getattr(edit, "rule_id", "") or ""))
        or _fallback_rule_text(edit)
    )


def _fallback_rule_text(edit: RuntimeEdit) -> str:
    edit_type = str(getattr(edit, "edit_type", "") or "").strip()
    if edit_type in {"spelling", "spelling_replace"}:
        return "Слово исправлено по нормативному написанию."
    if edit_type == "hyphen":
        return "Исправлено дефисное написание слова или сочетания."
    if edit_type == "split_join":
        return "Исправлено слитное или раздельное написание."
    if edit_type == "casing":
        return "Исправлено написание прописной или строчной буквы."
    if edit_type == "punctuation":
        return "Исправлена пунктуация."
    return "Орфографическая или пунктуационная правка."


def text_incremental_cache_key() -> str:
    return "default_text"


def build_incremental_summary(result: Any) -> dict[str, int]:
    return {
        "checked": int(getattr(result, "checked_segments", 0)),
        "reused": int(getattr(result, "reused_segments", 0)),
        "changed": int(getattr(result, "changed_segments", 0)),
    }


def format_incremental_status(summary: dict[str, int]) -> str:
    return (
        "Проверены только изменённые фрагменты: "
        f"заново проверено {summary['checked']}, "
        f"без повторной проверки {summary['reused']}, "
        f"новых или изменённых {summary['changed']}."
    )


def docx_cache_key(upload_name: Any) -> str:
    return _normalize_cache_key(upload_name, fallback="uploaded_docx")


def docx_correction_requested(upload_name: Any) -> bool:
    return bool(st.button("Исправить DOCX", type="primary", key=f"correct_docx_{docx_cache_key(upload_name)}"))


def build_docx_incremental_summary(result: Any) -> dict[str, int]:
    edits = getattr(result, "edits", []) or []
    return {
        "checked": int(getattr(result, "checked_paragraphs", 0)),
        "reused": int(getattr(result, "reused_paragraphs", 0)),
        "total_edits": len(edits),
    }


def render_docx_incremental_status(summary: dict[str, int] | None, *, enabled: bool) -> None:
    if not enabled or summary is None:
        return
    st.caption(format_docx_incremental_status(summary))


def format_docx_incremental_status(summary: dict[str, int]) -> str:
    return (
        "Повторно проверены только изменённые абзацы: "
        f"заново проверено {summary['checked']}, "
        f"без повторной проверки {summary['reused']}, "
        f"исправлений {summary['total_edits']}."
    )


def count_docx_text_blocks_for_gui(path: str | Path) -> int:
    document = Document(path)
    return sum(1 for paragraph in _iter_unique_docx_paragraphs(document) if paragraph.text.strip())


def _iter_unique_docx_paragraphs(document: Any) -> Any:
    seen: set[int] = set()
    for paragraph in document.paragraphs:
        paragraph_id = id(paragraph._p)
        if paragraph_id in seen:
            continue
        seen.add(paragraph_id)
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph_id = id(paragraph._p)
                    if paragraph_id in seen:
                        continue
                    seen.add(paragraph_id)
                    yield paragraph


def _import_streamlit() -> Any:
    import streamlit

    return streamlit


def _render_latest_text_correction() -> None:
    latest = st.session_state.get("latest_text_correction")
    if not latest:
        return

    source = str(latest["source"])
    result = latest["result"]
    st.text_area("Исправленный текст", value=result.corrected_text, height=180)
    summary = latest.get("incremental_summary")
    if summary:
        st.caption(format_incremental_status(summary))
    edits = list(getattr(result, "edits", []) or [])
    render_edit_summary(edits)
    st.markdown("Подсветка исправлений")
    st.markdown(render_highlighted_diff(source, result.corrected_text, edits), unsafe_allow_html=True)


def _latest_text_edits() -> list[RuntimeEdit]:
    latest = st.session_state.get("latest_text_correction")
    if not latest:
        return []
    result = latest["result"]
    return list(getattr(result, "edits", []) or [])


def _normalize_cache_key(value: Any, *, fallback: str) -> str:
    normalized = str(value or "").strip()
    return normalized or fallback


def render_highlighted_diff(source: str, corrected: str, edits: list[RuntimeEdit] | None = None) -> str:
    source_tokens = _diff_tokens(source)
    corrected_tokens = _diff_tokens(corrected)
    tooltips = _diff_tooltips(edits or [])
    chunks = _render_diff_chunks(source_tokens, corrected_tokens, tooltips)
    return f'{_DIFF_STYLE}<div class="diff-view">{"".join(chunks)}</div>'


def _diff_tooltips(edits: list[RuntimeEdit]) -> list[DiffTooltip]:
    return [
        DiffTooltip(
            source=str(getattr(edit, "source", "") or ""),
            replacement=str(getattr(edit, "replacement", "") or ""),
            text=_edit_tooltip_text(edit),
        )
        for edit in edits
    ]


def _edit_tooltip_text(edit: RuntimeEdit) -> str:
    return _user_facing_rule_text(edit)


def _render_diff_chunks(
    source_tokens: list[str],
    corrected_tokens: list[str],
    tooltips: list[DiffTooltip],
) -> list[str]:
    matcher = difflib.SequenceMatcher(
        a=_diff_match_keys(source_tokens),
        b=_diff_match_keys(corrected_tokens),
        autojunk=False,
    )
    return _render_diff_opcodes(source_tokens, corrected_tokens, matcher.get_opcodes(), tooltips)


def _render_diff_opcodes(
    source_tokens: list[str],
    corrected_tokens: list[str],
    opcodes: list[tuple[str, int, int, int, int]],
    tooltips: list[DiffTooltip],
) -> list[str]:
    chunks: list[str] = []
    for tag, source_start, source_end, corrected_start, corrected_end in opcodes:
        source_fragment = "".join(source_tokens[source_start:source_end])
        corrected_fragment = "".join(corrected_tokens[corrected_start:corrected_end])
        if tag == "equal":
            chunks.extend(
                _render_equal_tokens(
                    source_tokens[source_start:source_end],
                    corrected_tokens[corrected_start:corrected_end],
                    tooltips,
                )
            )
        elif tag == "delete":
            chunks.append(
                _diff_span(
                    "diff-delete",
                    source_fragment,
                    tooltip=_pop_delete_tooltip(tooltips, source_fragment),
                )
            )
        elif tag == "insert":
            chunks.append(
                _diff_span(
                    "diff-insert",
                    corrected_fragment,
                    tooltip=_pop_insert_tooltip(tooltips, corrected_fragment),
                )
            )
        elif tag == "replace":
            chunks.extend(
                _render_replace_tokens(
                    source_tokens[source_start:source_end],
                    corrected_tokens[corrected_start:corrected_end],
                    tooltips,
                )
            )
    return chunks


def _render_equal_tokens(
    source_tokens: list[str],
    corrected_tokens: list[str],
    tooltips: list[DiffTooltip],
) -> list[str]:
    chunks: list[str] = []
    for source_token, corrected_token in zip(source_tokens, corrected_tokens):
        if source_token == corrected_token:
            chunks.append(html.escape(corrected_token))
        else:
            chunks.append(_diff_span("diff-delete", source_token))
            chunks.append(
                _diff_span(
                    "diff-insert",
                    corrected_token,
                    tooltip=_pop_insert_tooltip(tooltips, corrected_token),
                )
            )
    return chunks


def _render_replace_tokens(
    source_tokens: list[str],
    corrected_tokens: list[str],
    tooltips: list[DiffTooltip],
) -> list[str]:
    whitespace_punctuation = _render_whitespace_punctuation_replacement(source_tokens, corrected_tokens, tooltips)
    if whitespace_punctuation is not None:
        return whitespace_punctuation

    matcher = difflib.SequenceMatcher(
        a=_diff_match_keys(source_tokens),
        b=_diff_match_keys(corrected_tokens),
        autojunk=False,
    )
    opcodes = matcher.get_opcodes()
    if len(opcodes) != 1 or opcodes[0][0] != "replace":
        return _render_diff_opcodes(source_tokens, corrected_tokens, opcodes, tooltips)

    return [
        _diff_span("diff-delete", "".join(source_tokens)),
        _diff_span(
            "diff-insert",
            "".join(corrected_tokens),
            tooltip=_pop_insert_tooltip(tooltips, "".join(corrected_tokens)),
        ),
    ]


def _render_whitespace_punctuation_replacement(
    source_tokens: list[str],
    corrected_tokens: list[str],
    tooltips: list[DiffTooltip],
) -> list[str] | None:
    source_text = "".join(source_tokens)
    corrected_text = "".join(corrected_tokens)
    if not source_text or not source_text.isspace() or not corrected_text:
        return None
    if any(char.isalnum() for char in corrected_text):
        return None
    if corrected_text.endswith(source_text):
        punctuation = corrected_text[: -len(source_text)]
        return (
            [
                _diff_span(
                    "diff-insert",
                    punctuation,
                    tooltip=_pop_insert_tooltip(tooltips, punctuation),
                ),
                html.escape(source_text),
            ]
            if punctuation
            else [html.escape(source_text)]
        )
    if corrected_text.startswith(source_text):
        punctuation = corrected_text[len(source_text) :]
        return (
            [
                html.escape(source_text),
                _diff_span(
                    "diff-insert",
                    punctuation,
                    tooltip=_pop_insert_tooltip(tooltips, punctuation),
                ),
            ]
            if punctuation
            else [html.escape(source_text)]
        )
    if not any(char.isspace() for char in corrected_text):
        return [_diff_span("diff-insert", corrected_text, tooltip=_pop_insert_tooltip(tooltips, corrected_text))]
    return None


def _pop_insert_tooltip(tooltips: list[DiffTooltip], value: str) -> str | None:
    return _pop_matching_tooltip(tooltips, value, field="replacement", require_empty_other=False)


def _pop_delete_tooltip(tooltips: list[DiffTooltip], value: str) -> str | None:
    return _pop_matching_tooltip(tooltips, value, field="source", require_empty_other=True)


def _pop_matching_tooltip(
    tooltips: list[DiffTooltip],
    value: str,
    *,
    field: str,
    require_empty_other: bool,
) -> str | None:
    if not value:
        return None
    value_key = _tooltip_match_key(value)
    if not value_key:
        return None

    candidates: list[tuple[int, DiffTooltip]] = []
    for index, tooltip in enumerate(tooltips):
        candidate = tooltip.replacement if field == "replacement" else tooltip.source
        other = tooltip.source if field == "replacement" else tooltip.replacement
        if require_empty_other and _tooltip_match_key(other):
            continue
        candidates.append((index, tooltip))
        if _tooltip_match_key(candidate) == value_key:
            return tooltips.pop(index).text

    if any(not char.isspace() and not char.isalnum() for char in value):
        for index, tooltip in candidates:
            candidate = tooltip.replacement if field == "replacement" else tooltip.source
            if value.strip() and value.strip() in candidate:
                return tooltips.pop(index).text
    if len(value_key) >= 3:
        for index, tooltip in candidates:
            candidate = tooltip.replacement if field == "replacement" else tooltip.source
            if value_key in _tooltip_token_keys(candidate):
                return tooltips.pop(index).text
    return None


def _tooltip_match_key(value: str) -> str:
    return str(value or "").strip().casefold()


def _tooltip_token_keys(value: str) -> set[str]:
    return {_tooltip_match_key(token) for token in _diff_tokens(value) if _tooltip_match_key(token)}


def _diff_span(css_class: str, value: str, *, tooltip: str | None = None) -> str:
    if not value:
        return ""
    escaped_value = html.escape(value)
    if not tooltip:
        return f'<span class="{css_class}">{escaped_value}</span>'
    escaped_tooltip = html.escape(tooltip, quote=True).replace("\n", "&#10;")
    return (
        f'<span class="{css_class} diff-tooltip" '
        f'data-tooltip="{escaped_tooltip}" title="{escaped_tooltip}" tabindex="0">{escaped_value}</span>'
    )


def _diff_match_keys(tokens: list[str]) -> list[str]:
    return [token.casefold() for token in tokens]


def _diff_tokens(text: str) -> list[str]:
    return re.findall(r"\s+|[А-Яа-яЁёA-Za-z0-9]+|[^\w\s]", text, flags=re.UNICODE)


_DIFF_STYLE = """
<style>
.diff-view {
    border: 1px solid #d8dee4;
    border-radius: 8px;
    padding: 0.85rem 1rem;
    background: #ffffff;
    color: #1f2328;
    line-height: 1.75;
    overflow: visible;
    white-space: pre-wrap;
    overflow-wrap: anywhere;
}
.diff-insert {
    background: #dafbe1;
    color: #116329;
    border-radius: 4px;
    padding: 0.08rem 0.18rem;
}
.diff-delete {
    background: #ffebe9;
    color: #82071e;
    border-radius: 4px;
    padding: 0.08rem 0.18rem;
    text-decoration: line-through;
}
.diff-tooltip {
    cursor: help;
    position: relative;
}
.diff-tooltip::after {
    background: #24292f;
    border-radius: 8px;
    bottom: calc(100% + 0.55rem);
    box-shadow: 0 8px 24px rgba(31, 35, 40, 0.18);
    color: #ffffff;
    content: attr(data-tooltip);
    font-size: 0.86rem;
    font-weight: 400;
    left: 50%;
    line-height: 1.35;
    max-width: min(420px, 80vw);
    opacity: 0;
    padding: 0.58rem 0.72rem;
    pointer-events: none;
    position: absolute;
    text-align: left;
    transform: translate(-50%, 0.2rem);
    transition: opacity 0.12s ease, transform 0.12s ease;
    visibility: hidden;
    white-space: pre-line;
    width: max-content;
    z-index: 1000;
}
.diff-tooltip:hover::after,
.diff-tooltip:focus::after {
    opacity: 1;
    transform: translate(-50%, 0);
    visibility: visible;
}
</style>
"""

_EDIT_PANEL_STYLE = """
<style>
.edit-summary {
    display: flex;
    flex-wrap: wrap;
    gap: 0.45rem;
    margin: 0.25rem 0 0.85rem;
}
.edit-summary span {
    border: 1px solid #d8dee4;
    border-radius: 999px;
    background: #f6f8fa;
    color: #24292f;
    font-size: 0.86rem;
    line-height: 1.2;
    padding: 0.32rem 0.58rem;
}
div[data-testid="stExpander"] {
    border-color: #d8dee4;
    border-radius: 8px;
}
div[data-testid="stExpander"] p {
    overflow-wrap: anywhere;
}
</style>
"""


if __name__ == "__main__":
    main()
